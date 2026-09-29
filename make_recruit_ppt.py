# -*- coding: utf-8 -*-
"""构建 AI Health Navigator 团队招募页 PPT（16:9，与 Demo 同视觉风格）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------- 配色（与 Demo 一致） ----------
TEAL   = RGBColor(0x0D, 0x6B, 0x5C)
TEAL_D = RGBColor(0x0A, 0x56, 0x4A)
MINT   = RGBColor(0x17, 0xB8, 0x97)
INK    = RGBColor(0x10, 0x22, 0x2B)
MUT    = RGBColor(0x4A, 0x5B, 0x64)
FAINT  = RGBColor(0x8A, 0x99, 0xA2)
BG     = RGBColor(0xF7, 0xFA, 0xF9)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
TINT   = RGBColor(0xE6, 0xF2, 0xEF)
LINE   = RGBColor(0xDD, 0xE6, 0xE3)

FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def shape(slide, st, l, t, w, h, fill=None, line=None, line_w=None, radius=None):
    shp = slide.shapes.add_shape(st, Inches(l), Inches(t), Inches(w), Inches(h))
    if radius is not None and st == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w or 1)
    try:
        shp.shadow.inherit = False
    except Exception:
        pass
    return shp


def shape_text(shp, text_, size, color, bold=True, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, font=FONT, line_spacing=1.0):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text_
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return shp


def add_text(slide, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             wrap=True, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get('align', align)
        para.space_before = Pt(p.get('sb', 0))
        para.space_after = Pt(p.get('sa', 0))
        para.line_spacing = p.get('ls', line_spacing)
        runs = p.get('runs')
        if runs is None:
            runs = [{'text': p.get('text', ''), 'size': p.get('size', 16),
                     'bold': p.get('bold', False), 'color': p.get('color', INK)}]
        for r in runs:
            run = para.add_run()
            run.text = r['text']
            f = run.font
            f.name = r.get('font', FONT)
            f.size = Pt(r.get('size', 16))
            f.bold = r.get('bold', False)
            f.color.rgb = r.get('color', INK)
    return box


def header(slide, kicker, title):
    shape(slide, MSO_SHAPE.RECTANGLE, 0.9, 0.82, 0.1, 0.5, fill=MINT)
    add_text(slide, 1.18, 0.76, 11.2, 0.3, [
        {'text': kicker, 'size': 13, 'bold': True, 'color': TEAL}])
    add_text(slide, 1.18, 1.08, 11.2, 0.85, [
        {'text': title, 'size': 31, 'bold': True, 'color': INK}])


def footer(slide, n):
    add_text(slide, 0.9, 7.08, 6, 0.3, [
        {'text': 'AI Health Navigator · 团队招募', 'size': 10, 'color': FAINT}])
    add_text(slide, 11.9, 7.08, 0.55, 0.3, [
        {'text': str(n), 'size': 11, 'bold': True, 'color': MUT}],
        align=PP_ALIGN.RIGHT)


def light_slide():
    s = prs.slides.add_slide(BLANK)
    bg(s, BG)
    return s


def steps_row(slide, y, items, x0=0.9, total_w=11.5, box_h=0.85,
              fill=TINT, txt_color=INK, size=14):
    n = len(items)
    gap = 0.4
    box_w = (total_w - gap * (n - 1)) / n
    x = x0
    for i, it in enumerate(items):
        bx = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h,
                   fill=fill, radius=0.5)
        shape_text(bx, it, size, txt_color, bold=True)
        if i < n - 1:
            add_text(slide, x + box_w, y, gap, box_h, [
                {'text': '→', 'size': 17, 'bold': True, 'color': MUT}],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += box_w + gap


def check_item(slide, y, mark, mark_color, title, desc):
    add_text(slide, 0.9, y, 0.5, 0.5, [
        {'text': mark, 'size': 20, 'bold': True, 'color': mark_color}])
    add_text(slide, 1.55, y, 5.4, 0.4, [
        {'text': title, 'size': 15.5, 'bold': True, 'color': INK}])
    add_text(slide, 1.55, y + 0.38, 5.4, 0.7, [
        {'text': desc, 'size': 12, 'color': MUT, 'ls': 1.2}])


# ============================================================
# 1. 封面
# ============================================================
s = prs.slides.add_slide(BLANK)
bg(s, TEAL)
mk = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 1.0, 0.72, 0.72, fill=MINT, radius=0.3)
shape_text(mk, "E", 28, TEAL_D, bold=True)
add_text(s, 1.85, 1.16, 9, 0.4, [
    {'text': 'AI HEALTH NAVIGATOR · 团队招募', 'size': 14, 'bold': True, 'color': MINT}])
add_text(s, 0.95, 2.05, 11.5, 1.1, [
    {'text': 'AI Health Navigator', 'size': 48, 'bold': True, 'color': WHITE}])
add_text(s, 0.95, 3.35, 11.5, 0.6, [
    {'text': '把日常健康的不确定性，连接到专业医疗。', 'size': 22, 'color': RGBColor(0xB8, 0xE4, 0xDA)}])
# 突出：已完成 Web Demo
bar = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 4.45, 8.3, 0.85, fill=MINT, radius=0.5)
shape_text(bar, '已完成一个可交互 Web Demo 原型 —— 就差你了', 17, TEAL_D, bold=True)
shape(s, MSO_SHAPE.RECTANGLE, 0.98, 5.85, 1.6, 0.045, fill=MINT)
add_text(s, 0.95, 6.15, 11.5, 0.6, [
    {'text': "We don't replace doctors. We connect everyday health uncertainty with professional care.",
     'size': 15, 'color': MINT, 'ls': 1.25}])

# ============================================================
# 2. 01 一句话介绍
# ============================================================
s = light_slide()
header(s, '01 · 一句话介绍', 'AI Health Navigator 是什么')
add_text(s, 0.9, 2.2, 11.5, 1.0, [
    {'text': '从身体出现异常，到真正见到医生之前，帮你管理整段健康旅程。',
     'size': 21, 'bold': True, 'color': INK, 'ls': 1.3}])
kw = [
    ('Self-Triage', '自我分诊'),
    ('Health Episode', '健康事件'),
    ('Continuous Follow-up', '持续追踪'),
    ('Doctor Handoff', '就医摘要'),
]
x = 0.9
for en, zh in kw:
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.5, 2.75, 1.35, fill=WHITE,
              line=LINE, line_w=1.2, radius=0.12)
    add_text(s, x + 0.22, 3.72, 2.3, 0.4, [
        {'text': en, 'size': 13, 'bold': True, 'color': TEAL}])
    add_text(s, x + 0.22, 4.15, 2.3, 0.4, [
        {'text': zh, 'size': 15.5, 'bold': True, 'color': INK}])
    x += 2.92
c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 5.4, 11.5, 0.8, fill=TINT, radius=0.3)
add_text(s, 1.25, 5.4, 10.8, 0.8, [
    {'runs': [
        {'text': '已有可交互 Web Demo（虚构数据 · 不替代医疗建议）', 'size': 14, 'bold': True, 'color': TEAL},
        {'text': '   首页 / 结构化问询 / Episode 时间线 / 症状趋势 / 就医摘要', 'size': 13, 'color': MUT},
    ]}], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)

# ============================================================
# 3. 02 我们发现的问题
# ============================================================
s = light_slide()
header(s, '02 · 问题', '今天的不舒服，卡在哪一步')
add_text(s, 0.9, 2.25, 5.3, 0.9, [
    {'text': '用户真正困惑的，往往不是「这是什么病」，而是：', 'size': 17, 'bold': True,
     'color': INK, 'ls': 1.25}])
qs = ['我现在该怎么办？', '需要去医院吗？', '什么时候不能再等？']
for i, q in enumerate(qs):
    y = 3.2 + i * 0.78
    shape(s, MSO_SHAPE.RECTANGLE, 0.9, y, 0.09, 0.62, fill=MINT)
    add_text(s, 1.15, y, 5.0, 0.62, [
        {'text': q, 'size': 20, 'bold': True, 'color': TEAL}], anchor=MSO_ANCHOR.MIDDLE)
# 右侧：搜索循环
c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.7, 2.25, 5.7, 4.15, fill=WHITE,
          line=LINE, line_w=1.2, radius=0.08)
add_text(s, 7.1, 2.6, 4.9, 0.5, [
    {'text': '现有的「搜索循环」', 'size': 16, 'bold': True, 'color': INK}])
loop = ['反复搜索 / 问 AI', '得到碎片化答案', '越来越不确定、更焦虑', '继续搜索']
for i, t in enumerate(loop):
    y = 3.3 + i * 0.72
    add_text(s, 7.1, y, 4.6, 0.4, [
        {'text': t, 'size': 14.5, 'bold': True, 'color': MUT}])
    if i < len(loop) - 1:
        add_text(s, 7.25, y + 0.42, 0.5, 0.28, [
            {'text': '↓', 'size': 15, 'bold': True, 'color': FAINT}])
add_text(s, 7.25, 3.3 + 4 * 0.72 + 0.02, 4.6, 0.35, [
    {'text': '↻ 又回到原点', 'size': 13, 'bold': True, 'color': TEAL}])
footer(s, 3)

# ============================================================
# 4. 03 我们的 Solution
# ============================================================
s = light_slide()
header(s, '03 · Solution', '一次 Health Episode 管到底')
add_text(s, 0.9, 2.35, 11.5, 0.6, [
    {'text': '不是搜索引擎式的症状罗列，而是一条有始有终的线索。', 'size': 17, 'bold': True,
     'color': INK, 'ls': 1.2}])
steps_row(s, 3.3, ['描述不适', '结构化分诊', '建立健康事件', '持续追踪', '就医摘要'],
          fill=WHITE, txt_color=TEAL, size=13.5, box_h=0.95)
add_text(s, 0.9, 4.6, 11.5, 0.4, [
    {'text': '每一步，都有据可循；每一次变化，都被记录。', 'size': 13.5, 'color': MUT}])
c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 5.35, 11.5, 1.15, fill=TINT, radius=0.14)
add_text(s, 1.35, 5.6, 10.6, 0.5, [
    {'text': "We don't replace doctors.", 'size': 21, 'bold': True, 'color': TEAL}])
add_text(s, 1.35, 6.05, 10.6, 0.4, [
    {'text': 'We connect everyday health uncertainty with professional care.',
     'size': 14, 'color': MUT}])
footer(s, 4)

# ============================================================
# 5. 04 为什么现在值得做
# ============================================================
s = light_slide()
header(s, '04 · 为什么现在', '三个趋势在此交汇')
cards = [
    ('AI', 'LLM 让多轮、可解释的对话式问询真正可用，不再是关键词匹配。'),
    ('Data', '个人健康数据（可穿戴 / 手动记录）终于可以被持续追踪与解释。'),
    ('Healthcare Gap', '日常的不确定与专业医疗之间，缺一个可信的连接层。'),
]
x = 0.9
for en, desc in cards:
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.55, 3.7, 3.3, fill=WHITE,
              line=LINE, line_w=1.2, radius=0.1)
    add_text(s, x + 0.3, 3.0, 3.1, 0.6, [
        {'text': en, 'size': 20, 'bold': True, 'color': TEAL}])
    add_text(s, x + 0.3, 3.7, 3.1, 1.7, [
        {'text': desc, 'size': 14, 'color': MUT, 'ls': 1.35}])
    x += 3.9
add_text(s, 0.9, 6.2, 11.5, 0.5, [
    {'text': '现在做，不是「能做什么」，而是「正好能做」。', 'size': 14, 'bold': True, 'color': INK}])
footer(s, 5)

# ============================================================
# 6. 05 我们已经做到哪里
# ============================================================
s = light_slide()
header(s, '05 · 进度', '已经不是一张 PPT')
add_text(s, 0.9, 2.3, 5.6, 0.4, [
    {'text': '已完成', 'size': 13, 'bold': True, 'color': TEAL}])
check_item(s, 2.75, '✓', MINT, '可交互 Web Demo',
           '首页 / 结构化问询 / Episode 时间线 / 症状趋势 / 就医摘要')
check_item(s, 4.05, '✓', MINT, '项目 Demo PPT + 演示动画',
           '产品定位、核心用户流程、商业模式')
check_item(s, 5.35, '✓', MINT, '产品定位与核心流程',
           'Self-Triage · Health Episode · Follow-up · Handoff')
add_text(s, 0.9, 6.4, 5.6, 0.4, [
    {'text': '下一步：真实用户访谈 → 快速迭代 → 验证假设', 'size': 13.5, 'bold': True,
     'color': INK}])
# 右侧：二维码占位
qr = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.6, 2.5, 4.7, 4.2, fill=WHITE,
           line=LINE, line_w=1.4, radius=0.06)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.85, 3.0, 2.2, 2.2, fill=BG, line=LINE, line_w=1.0, radius=0.06)
add_text(s, 8.85, 5.3, 2.2, 0.4, [
    {'text': 'Demo 二维码', 'size': 13, 'bold': True, 'color': TEAL}], align=PP_ALIGN.CENTER)
add_text(s, 8.0, 5.8, 3.9, 0.7, [
    {'text': '扫码体验可交互 Demo', 'size': 15, 'bold': True, 'color': INK}], align=PP_ALIGN.CENTER)
footer(s, 6)

# ============================================================
# 7. 06 我们正在寻找谁
# ============================================================
s = light_slide()
header(s, '06 · 找谁', '5–6 人的小团队')
roles = [
    ('Product', '产品', '× 1'),
    ('AI / LLM', 'AI', '× 1–2'),
    ('Full-stack', '全栈', '× 1'),
    ('Medical', '医学', '× 1'),
    ('Design', '设计', '× 1'),
]
x = 0.9
cw = 2.22
for en, zh, num in roles:
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.6, cw, 2.9, fill=WHITE,
              line=LINE, line_w=1.2, radius=0.12)
    shape_text(shape(s, MSO_SHAPE.OVAL, x + 0.7, 3.0, 0.82, 0.82, fill=TINT),
               num, 15, TEAL, bold=True)
    add_text(s, x + 0.18, 4.1, cw - 0.36, 0.45, [
        {'text': zh, 'size': 17, 'bold': True, 'color': INK}], align=PP_ALIGN.CENTER)
    add_text(s, x + 0.18, 4.6, cw - 0.36, 0.4, [
        {'text': en, 'size': 11.5, 'color': MUT}], align=PP_ALIGN.CENTER)
    x += cw + 0.1
add_text(s, 0.9, 5.9, 11.5, 0.5, [
    {'text': '一个人扛不动一个产品；我们互补，一起把它跑起来。', 'size': 14, 'bold': True, 'color': INK}])
footer(s, 7)

# ============================================================
# 8. 07 你会真正参与什么
# ============================================================
s = light_slide()
header(s, '07 · 你会做什么', '不是查资料 + 做 PPT')
add_text(s, 0.9, 2.35, 11.5, 0.5, [
    {'text': '从第一天起，就参与真实的判断与执行：', 'size': 17, 'bold': True, 'color': INK}])
steps = ['用户访谈', '快速迭代', 'Demo', 'MVP', '真实用户', '验证假设']
x = 0.9
cw = 1.75
for i, t in enumerate(steps):
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.4, cw, 1.0, fill=WHITE,
              line=LINE, line_w=1.2, radius=0.5)
    shape_text(c, t, 14, TEAL, bold=True)
    if i < len(steps) - 1:
        add_text(s, x + cw, 3.4, 0.4, 1.0, [
            {'text': '→', 'size': 17, 'bold': True, 'color': MUT}],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    x += cw + 0.4
add_text(s, 0.9, 4.9, 11.5, 1.3, [
    {'text': '每个阶段都动手：写代码 / 做访谈 / 画原型 / 看数据，而不是旁观。',
     'size': 15, 'color': MUT, 'ls': 1.3, 'sa': 6},
    {'text': 'Demo → MVP → 真实用户 → 验证假设', 'size': 16, 'bold': True, 'color': TEAL},
])
footer(s, 8)

# ============================================================
# 9. 08 希望找到怎样的队友
# ============================================================
s = light_slide()
header(s, '08 · 队友', '我们希望你是这样的人')
traits = [
    ('Builder mindset', '动手造，而不是只讨论'),
    ('Curiosity', '对「为什么」保持好奇'),
    ('Ownership', '把自己负责的事做到位'),
    ('Fast learning', '边做边学，快速迭代'),
]
x = 0.9
for en, zh in traits:
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.5, 2.75, 2.2, fill=WHITE,
              line=LINE, line_w=1.2, radius=0.1)
    add_text(s, x + 0.22, 2.75, 2.3, 0.5, [
        {'text': en, 'size': 15, 'bold': True, 'color': TEAL}])
    add_text(s, x + 0.22, 3.45, 2.3, 0.9, [
        {'text': zh, 'size': 14, 'color': MUT, 'ls': 1.25}])
    x += 2.92
c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 5.2, 11.5, 0.85, fill=TINT, radius=0.3)
add_text(s, 1.25, 5.2, 10.8, 0.85, [
    {'text': '年级、经历不重要；重要的是，你想不想一起把这件事做成。', 'size': 15.5,
     'bold': True, 'color': TEAL}], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 9)

# ============================================================
# 10. 09 为什么加入
# ============================================================
s = light_slide()
header(s, '09 · 为什么加入', '你能带走什么')
items = [
    ('0 → 1 的完整经历', '从想法、Demo 到真实用户，全程参与，不是打杂。'),
    ('跨背景协作', '产品 × AI × 医学 × 设计，学会与不同的人一起做产品。'),
    ('可长期的方向', '可发展为创业、科研或竞赛项目，不是一次性作业。'),
]
x = 0.9
for t, d in items:
    c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.5, 3.7, 3.1, fill=WHITE,
              line=LINE, line_w=1.2, radius=0.1)
    add_text(s, x + 0.28, 2.85, 3.15, 0.55, [
        {'text': t, 'size': 17, 'bold': True, 'color': TEAL}])
    add_text(s, x + 0.28, 3.5, 3.15, 1.5, [
        {'text': d, 'size': 13.5, 'color': MUT, 'ls': 1.35}])
    x += 3.9
footer(s, 10)

# ============================================================
# 11. 10 招募方式
# ============================================================
s = light_slide()
header(s, '10 · Apply', '如何加入')
add_text(s, 0.9, 2.3, 11.5, 0.5, [
    {'runs': [
        {'text': '团队规模：', 'size': 16, 'bold': True, 'color': INK},
        {'text': '5–6 人，小而精。', 'size': 16, 'color': INK},
    ]}])
add_text(s, 0.9, 3.0, 11.5, 0.4, [
    {'text': '申请方式：', 'size': 16, 'bold': True, 'color': INK}])
add_text(s, 0.9, 3.45, 11.5, 0.5, [
    {'text': '填写报名表（含一道开放题），我们会尽快与你聊一聊。', 'size': 14.5, 'color': MUT}])
c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 4.35, 11.5, 1.6, fill=TINT, radius=0.12)
add_text(s, 1.35, 4.6, 10.6, 0.4, [
    {'text': '开放题', 'size': 12.5, 'bold': True, 'color': TEAL}])
add_text(s, 1.35, 5.0, 10.6, 0.8, [
    {'text': '“如果由你设计这个产品，你最想解决的一个问题是什么？”', 'size': 20, 'bold': True,
     'color': INK, 'ls': 1.2}])
add_text(s, 0.9, 6.35, 11.5, 0.5, [
    {'text': '期待与你一起，从 0 做到 1。', 'size': 14, 'bold': True, 'color': TEAL}])
footer(s, 11)

out = r"C:\Users\29383\Documents\Codex\2026-09-29\referenced-chatgpt-conversation-this-is-an-2\outputs\episode\AI-Health-Navigator-Recruit.pptx"
prs.save(out)
print("saved:", out)
print("slides:", len(prs.slides._sldIdLst))
